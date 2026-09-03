"""
In-app notifications + email delivery + personal timetable export.

Email is entirely optional: if SMTP_HOST isn't configured, send_email() just
records that it was skipped instead of failing the request — a deployment
without SMTP set up yet should still be able to publish rosters and use
in-app notifications.
"""
import io
import os
import smtplib
from datetime import date, datetime
from email.message import EmailMessage
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlmodel import Session, select

from .database import get_session
from .models import Notification, Employee, RosterRun, Assignment
from .auth import get_current_user, require_role
from .scheduler import date_range

router = APIRouter(prefix="/api", tags=["notifications"])

any_role = require_role("admin", "call_center_tl", "drivers_manager", "phtl")

SMTP_HOST = os.environ.get("SMTP_HOST")
SMTP_PORT = int(os.environ.get("SMTP_PORT", "587"))
SMTP_USER = os.environ.get("SMTP_USER")
SMTP_PASSWORD = os.environ.get("SMTP_PASSWORD")
SMTP_FROM = os.environ.get("SMTP_FROM", SMTP_USER or "no-reply@samu-roster.local")


def send_email(to_email: str, subject: str, body: str,
                attachment_bytes: Optional[bytes] = None, attachment_name: Optional[str] = None) -> bool:
    if not SMTP_HOST or not to_email:
        return False
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = SMTP_FROM
    msg["To"] = to_email
    msg.set_content(body)
    if attachment_bytes:
        msg.add_attachment(attachment_bytes, maintype="application",
                            subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            filename=attachment_name or "timetable.xlsx")
    try:
        with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=10) as server:
            server.starttls()
            if SMTP_USER:
                server.login(SMTP_USER, SMTP_PASSWORD or "")
            server.send_message(msg)
        return True
    except Exception:
        return False


def build_employee_timetable(session: Session, run: RosterRun, employee: Employee) -> bytes:
    days = date_range(run.start_date, run.end_date)
    rows = session.exec(select(Assignment).where(
        Assignment.run_id == run.id, Assignment.employee_id == employee.id
    )).all()
    by_day = {a.day: a.code for a in rows}

    wb = Workbook()
    ws = wb.active
    ws.title = "Timetable"
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    ws.cell(row=1, column=1, value=employee.name).font = header_font
    ws.cell(row=1, column=1).fill = header_fill
    ws.cell(row=2, column=1, value="Day").font = Font(bold=True)
    ws.cell(row=2, column=2, value="Assignment").font = Font(bold=True)
    for i, d in enumerate(days, start=3):
        ws.cell(row=i, column=1, value=d.isoformat())
        ws.cell(row=i, column=2, value=by_day.get(d, ""))
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 14

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def notify_employee(session: Session, employee_id: int, title: str, body: str,
                     related_type: Optional[str] = None, related_id: Optional[int] = None,
                     attach_run: Optional[RosterRun] = None):
    """Create the in-app record and, if the employee has an email and SMTP is
    configured, also email them (with their personal timetable attached when
    attach_run is given)."""
    employee = session.get(Employee, employee_id)
    channel = "in_app"
    sent_at = None
    if employee and employee.email:
        attachment_bytes, attachment_name = None, None
        if attach_run:
            attachment_bytes = build_employee_timetable(session, attach_run, employee)
            attachment_name = f"timetable_{attach_run.start_date.isoformat()}_{attach_run.end_date.isoformat()}.xlsx"
        if send_email(employee.email, title, body, attachment_bytes, attachment_name):
            channel = "email"
            sent_at = datetime.utcnow()
    notification = Notification(
        employee_id=employee_id, title=title, body=body,
        related_type=related_type, related_id=related_id,
        channel=channel, sent_at=sent_at,
    )
    session.add(notification)


@router.get("/notifications")
def list_notifications(session: Session = Depends(get_session), user=Depends(get_current_user)):
    if not user.employee_id:
        return []
    rows = session.exec(
        select(Notification).where(Notification.employee_id == user.employee_id)
        .order_by(Notification.id.desc())
    ).all()
    return rows


@router.put("/notifications/{notification_id}/read")
def mark_read(notification_id: int, session: Session = Depends(get_session), user=Depends(get_current_user)):
    note = session.get(Notification, notification_id)
    if not note or note.employee_id != user.employee_id:
        raise HTTPException(404, "Notification not found")
    note.read_at = datetime.utcnow()
    session.add(note)
    session.commit()
    session.refresh(note)
    return note


@router.get("/roster/{run_id}/export/{employee_id}")
def export_employee_timetable(run_id: int, employee_id: int,
                               session: Session = Depends(get_session), _=Depends(any_role)):
    from fastapi.responses import StreamingResponse
    run = session.get(RosterRun, run_id)
    if not run:
        raise HTTPException(404, "Roster run not found")
    employee = session.get(Employee, employee_id)
    if not employee:
        raise HTTPException(404, "Employee not found")
    content = build_employee_timetable(session, run, employee)
    filename = f"timetable_{employee.name.replace(' ', '_')}_{run.start_date.isoformat()}_{run.end_date.isoformat()}.xlsx"
    return StreamingResponse(
        io.BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
